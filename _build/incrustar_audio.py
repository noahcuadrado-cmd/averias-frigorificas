#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Incrusta un MP3 por diapositiva en un PPTX con AUTOPLAY y AVANCE AUTOMATICO.
Uso: python3 incrustar_audio.py <deck.pptx> <carpeta_mp3> <prefijo tXXvN> <salida.pptx> <carpeta_wav_para_duracion>
- El MP3 de la diapositiva i se llama <prefijo>_<i:03d>.mp3
- Diapositivas sin audio: avance por defecto (4 s).
"""
import sys, os, glob, wave, contextlib
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

def dur_wav(path):
    with contextlib.closing(wave.open(path,'rb')) as w:
        return w.getnframes()/float(w.getframerate())

TIMING = '''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetClass="mediaCall" presetID="0" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="indefinite" fill="hold"/><p:tgtEl><p:spid>%SPID%</p:spid></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
</p:seq>
<p:audio><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spid>%SPID%</p:spid></p:tgtEl></p:cMediaNode></p:audio>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''

def set_advance(slide, ms):
    sld=slide._element
    # los <p:transition> (cubo) estan dentro de mc:AlternateContent; poner advTm en todos
    for tr in sld.iter(qn('p:transition')):
        tr.set('advTm', str(int(ms))); tr.set('advClick','0')

def add_timing(slide, spid):
    sld=slide._element
    # quitar timing previo
    for old in sld.findall(qn('p:timing')): sld.remove(old)
    node=parse_xml(TIMING.replace('%SPID%', str(spid)))
    sld.append(node)  # timing va al final del sld

def main():
    deck, mp3dir, pref, out, wavdir = sys.argv[1:6]
    prs=Presentation(deck)
    slides=list(prs.slides)
    n_audio=0
    for i,slide in enumerate(slides,1):
        mp3=os.path.join(mp3dir, f"{pref}_{i:03d}.mp3")
        wav=os.path.join(wavdir, f"{pref}_{i:03d}.wav")
        if os.path.exists(mp3):
            d=dur_wav(wav) if os.path.exists(wav) else 20.0
            mv=slide.shapes.add_movie(mp3, Emu(-3000000), Emu(-3000000), Emu(900000), Emu(900000), mime_type='audio/mpeg')
            spid=mv.shape_id
            # videoFile -> audioFile (icono de audio)
            for vf in mv._element.iter(qn('a:videoFile')): vf.tag=qn('a:audioFile')
            add_timing(slide, spid)
            set_advance(slide, round((d+0.7)*1000))
            n_audio+=1
        else:
            set_advance(slide, 4000)  # sin audio -> 4 s
    prs.save(out)
    print(f"Incrustados {n_audio} audios en {len(slides)} diapositivas -> {out}")

if __name__=="__main__":
    main()
