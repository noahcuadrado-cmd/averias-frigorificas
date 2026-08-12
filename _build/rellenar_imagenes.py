#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rellena las IMAGENES realistas de los PowerPoint del Curso Gas B.
SE EJECUTA EN TU ORDENADOR (que si tiene Internet), NO en la nube.

Cada diapositiva de contenido lleva un marco con una etiqueta oculta
"IMGQ:<termino de busqueda>". Este script busca una foto realista para ese
termino (API gratuita de Pexels), la descarga y la incrusta en el marco.

--- REQUISITOS (una sola vez) ---
1) Instala Python 3 en tu PC (python.org). En Windows marca "Add to PATH".
2) Abre una terminal (PowerShell) y ejecuta:
       pip install python-pptx requests
3) Consigue una clave GRATUITA de Pexels:
       https://www.pexels.com/api/  ->  "Get Started"  ->  copia tu API Key
   (Alternativa Unsplash: pon USAR="unsplash" y tu Access Key.)

--- USO ---
   python rellenar_imagenes.py  "C:\\ruta\\a\\Curso Gas B"  TU_CLAVE_PEXELS

   - El primer argumento es la carpeta con los .pptx (busca en subcarpetas).
   - El segundo es tu clave.
   - Genera copias "<nombre> (con imagenes).pptx" (no toca los originales).
   - Si una diapositiva ya tenia foto, la respeta.

Consejo: si una foto no te gusta, cambia el termino en las notas/etiqueta y
vuelve a pasarlo, o ejecuta con --idx 2 para coger la 2a opcion de busqueda.
"""
import sys, os, io, time, glob, argparse, urllib.request, urllib.parse, json

USAR = "pexels"   # "pexels" o "unsplash"

def buscar_pexels(q, key, idx=1, intento=0):
    url = "https://api.pexels.com/v1/search?" + urllib.parse.urlencode(
        {"query": q, "per_page": max(idx,1), "orientation": "landscape", "locale": "es-ES"})
    req = urllib.request.Request(url, headers={"Authorization": key})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            data = json.load(r)
        fotos = data.get("photos", [])
        if not fotos: return None
        f = fotos[min(idx-1, len(fotos)-1)]
        return f["src"].get("large") or f["src"].get("original")
    except Exception as e:
        if intento < 3:
            time.sleep(2*(intento+1)); return buscar_pexels(q, key, idx, intento+1)
        print("   ! error buscando:", e); return None

def buscar_unsplash(q, key, idx=1, intento=0):
    url = "https://api.unsplash.com/search/photos?" + urllib.parse.urlencode(
        {"query": q, "per_page": max(idx,1), "orientation": "landscape"})
    req = urllib.request.Request(url, headers={"Authorization": "Client-ID "+key})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            data = json.load(r)
        res = data.get("results", [])
        if not res: return None
        return res[min(idx-1, len(res)-1)]["urls"].get("regular")
    except Exception as e:
        if intento < 3:
            time.sleep(2*(intento+1)); return buscar_unsplash(q, key, idx, intento+1)
        print("   ! error buscando:", e); return None

def descargar(url):
    with urllib.request.urlopen(url, timeout=60) as r:
        return io.BytesIO(r.read())

def procesar(path, key, idx, cache):
    from pptx import Presentation
    prs = Presentation(path)
    n_ok = 0
    for si, slide in enumerate(prs.slides, 1):
        # buscar el marco-placeholder por su descr "IMGQ:"
        ph = None
        for sh in slide.shapes:
            try:
                cnv = sh._element.xpath('.//p:cNvPr')[0]
                d = cnv.get('descr') or ""
            except Exception:
                d = ""
            if d.startswith("IMGQ:"):
                ph = sh; q = d[5:].strip(); break
        if ph is None:
            continue
        L, T, Wd, Ht = ph.left, ph.top, ph.width, ph.height
        if q not in cache:
            url = (buscar_pexels if USAR == "pexels" else buscar_unsplash)(q, key, idx)
            cache[q] = url
            time.sleep(0.5)   # cortesia con la API
        url = cache[q]
        if not url:
            print(f"   slide {si}: sin resultado para '{q}' (se deja el marco)")
            continue
        try:
            bio = descargar(url)
            slide.shapes.add_picture(bio, L, T, width=Wd, height=Ht)
            # quitar el marco/etiqueta (la foto ya lo cubre)
            ph._element.getparent().remove(ph._element)
            n_ok += 1
        except Exception as e:
            print(f"   slide {si}: error incrustando '{q}':", e)
    salida = os.path.splitext(path)[0] + " (con imagenes).pptx"
    prs.save(salida)
    print(f"  -> {os.path.basename(salida)}  ({n_ok} imagenes)")
    return n_ok

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("carpeta"); ap.add_argument("clave")
    ap.add_argument("--idx", type=int, default=1, help="que resultado de busqueda usar (1=primero)")
    a = ap.parse_args()
    pptx = [p for p in glob.glob(os.path.join(a.carpeta, "**", "*.pptx"), recursive=True)
            if "(con imagenes)" not in p]
    if not pptx:
        print("No se han encontrado .pptx en", a.carpeta); return
    print(f"{len(pptx)} presentaciones. Fuente de fotos: {USAR}.")
    cache = {}; tot = 0
    for p in sorted(pptx):
        print("*", os.path.relpath(p, a.carpeta))
        tot += procesar(p, a.clave, a.idx, cache)
    print(f"\nHecho. {tot} imagenes incrustadas en total.")

if __name__ == "__main__":
    main()
